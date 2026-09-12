"""Convert WRITEUP_DOC.md into a .docx for import into Google Docs.

Handles: # headings, **bold** lead-ins, bullet lists, pipe tables, italic captions,
inline code, and the [FIGURE 1: path] marker (inserted as an image, 6.5in wide).
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

SRC = Path("WRITEUP_DOC.md")
OUT = Path("WRITEUP_DOC.docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(11)
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)


def add_runs(par, text):
    """Add text with **bold**, *italic*, `code` inline formatting."""
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for t in tokens:
        if not t:
            continue
        if t.startswith("**"):
            r = par.add_run(t[2:-2]); r.bold = True
        elif t.startswith("*"):
            r = par.add_run(t[1:-1]); r.italic = True
        elif t.startswith("`"):
            r = par.add_run(t[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(10)
        else:
            par.add_run(t)


def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for c in cells if not all(re.fullmatch(r":?-+:?", x) for x in c)]
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, txt)
            for r in p.runs:
                r.font.size = Pt(9.5)
                if i == 0:
                    r.bold = True
    doc.add_paragraph()


lines = SRC.read_text().splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith("|"):
        block = []
        while i < len(lines) and lines[i].startswith("|"):
            block.append(lines[i]); i += 1
        add_table(block)
        continue
    m = re.match(r"^(#+)\s+(.*)", line)
    if m:
        level = len(m.group(1))
        doc.add_heading(m.group(2), level=min(level, 3) if level > 1 else 0)
        i += 1
        continue
    fm = re.match(r"^\[FIGURE 1: (.+)\]", line)
    if fm:
        doc.add_picture(fm.group(1).strip(), width=Inches(6.5))
        i += 1
        continue
    if line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, line[2:])
        i += 1
        continue
    if line.strip() == "":
        i += 1
        continue
    p = doc.add_paragraph()
    add_runs(p, line)
    i += 1

doc.save(OUT)
print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
